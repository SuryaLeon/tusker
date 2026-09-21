from pathlib import Path

import pandas as pd
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

def _add_schedule_comments(
    worksheet,
    schedule_df,
):
    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    score_column = headers.get(
        "Score"
    )

    if score_column is None:
        return

    for index, row in (
        schedule_df.iterrows()
    ):
        excel_row = index + 2

        score = row.get(
            "Score"
        )

        reason = row.get(
            "Reason",
            "",
        )

        poc1 = row.get(
            "POC1",
            "",
        )

        poc2 = row.get(
            "POC2",
            "",
        )

        comment_text = (
            f"Selected pair: "
            f"{poc1} + {poc2}\n"
            f"Score: {score}\n\n"
            f"Reason:\n{reason}"
        )

        cell = worksheet.cell(
            row=excel_row,
            column=score_column,
        )

        cell.comment = Comment(
            comment_text,
            "POC Scheduler",
        )

def _format_worksheet(worksheet):
    """
    Apply lightweight human-friendly Excel formatting.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1\]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for column_cells in worksheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            try:
                value = str(
                    cell.value
                    if cell.value is not None
                    else ""
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            except Exception:
                pass

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 10),
            60,
        )


def export_schedule(
    schedule_df,
    rankings_df,
    pair_statistics_df,
    poc_statistics_df,
    output_path,
):
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        schedule_df.to_excel(
            writer,
            sheet_name="Schedule",
            index=False,
        )

        rankings_df.to_excel(
            writer,
            sheet_name="Candidate Rankings",
            index=False,
        )

        pair_statistics_df.to_excel(
            writer,
            sheet_name="Pair Statistics",
            index=False,
        )

        poc_statistics_df.to_excel(
            writer,
            sheet_name="POC Statistics",
            index=False,
        )

        for worksheet in writer.book.worksheets:
            _format_worksheet(
                worksheet
            )

    return output_path